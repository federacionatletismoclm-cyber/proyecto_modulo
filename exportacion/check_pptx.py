import os, io, sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

base = os.path.dirname(os.path.abspath(__file__))
prs = Presentation(os.path.join(base, 'Presentacion_Mejora_Pistas_Multideporte.pptx'))
PX = 6350
FD = r'C:\Windows\Fonts'
def font(name, bold, size_px):
    cand = {('Arial Narrow', False): 'arialn.ttf', ('Arial Narrow', True): 'arialnb.ttf',
            ('Segoe UI Symbol', False): 'seguisym.ttf', ('Segoe UI Symbol', True): 'seguisym.ttf', ('Calibri', False): 'calibri.ttf', ('Calibri', True): 'calibrib.ttf'}[(name, bool(bold))]
    return ImageFont.truetype(os.path.join(FD, cand), int(size_px))

def wrap(text_runs, width):
    """text_runs: list[(text,bold,fontname,size_px)] -> lista de líneas (alto en px)"""
    words = []
    for t, b, n, s in text_runs:
        for w in t.replace('\n', ' \n ').split(' '):
            if w == '': continue
            words.append((w, b, n, s))
    lines, cur, curw, maxs = [], [], 0, 0
    d = ImageDraw.Draw(Image.new('RGB', (10, 10)))
    for w, b, n, s in words:
        f = font(n, b, s); ww = d.textlength(w + ' ', font=f)
        if cur and curw + ww > width + 1:
            lines.append((cur, maxs)); cur, curw, maxs = [], 0, 0
        cur.append((w, b, n, s)); curw += ww; maxs = max(maxs, s)
    if cur: lines.append((cur, maxs))
    return lines

problems = []
sheet_w, sheet_h = 640, 360
sheet = Image.new('RGB', (sheet_w * 3, sheet_h * 4), 'white')
for si, slide in enumerate(prs.slides, 1):
    canvas = Image.new('RGB', (1920, 1080), 'white'); dr = ImageDraw.Draw(canvas, 'RGBA')
    bgc = slide.background.fill.fore_color.rgb
    dr.rectangle([0, 0, 1920, 1080], fill=tuple(bgc))
    for sh in slide.shapes:
        x, y, w, h = [v / PX for v in (sh.left, sh.top, sh.width, sh.height)]
        if x < -1 or y < -1 or x + w > 1921 or y + h > 1081:
            problems.append(f'Diap {si}: forma fuera del lienzo ({sh.shape_type}, x={x:.0f}, y={y:.0f}, w={w:.0f}, h={h:.0f})')
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            im = Image.open(io.BytesIO(sh.image.blob)).convert('RGB')
            iw, ih = im.size
            l, r, t, b = sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom
            im = im.crop((int(l * iw), int(t * ih), int(iw - r * iw), int(ih - b * ih))).resize((int(w), int(h)))
            canvas.paste(im, (int(x), int(y)))
            dr = ImageDraw.Draw(canvas, 'RGBA')
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            sp = sh._element.spPr
            from pptx.oxml.ns import qn
            if sp.find(qn('a:gradFill')) is not None:
                gs = sp.findall('.//' + qn('a:gs'))
                stops = [(int(g.get('pos')) / 100000, int(g.find(qn('a:srgbClr')).find(qn('a:alpha')).get('val')) / 100000) for g in gs]
                ang = int(sp.find('.//' + qn('a:lin')).get('ang')) / 60000
                ov = Image.new('RGBA', (int(w), int(h)))
                px = ov.load()
                for yy in range(int(h)):
                    for xx in range(int(w)):
                        t = xx / w if ang == 0 else yy / h
                        for (p0, a0), (p1, a1) in zip(stops, stops[1:]):
                            if p0 <= t <= p1:
                                a = a0 + (a1 - a0) * ((t - p0) / (p1 - p0 or 1)); break
                        px[xx, yy] = (16, 36, 58, int(255 * a))
                canvas.paste(ov, (int(x), int(y)), ov); dr = ImageDraw.Draw(canvas, 'RGBA')
            else:
                fill = None
                if sh.fill.type == 1:
                    c = sh.fill.fore_color.rgb; a = 255
                    al = sp.find('.//' + qn('a:alpha'))
                    if al is not None: a = int(int(al.get('val')) / 100000 * 255)
                    fill = (c[0], c[1], c[2], a)
                outline = None
                if sh.line.fill.type == 1: outline = tuple(sh.line.color.rgb)
                kind = sh.auto_shape_type
                rad = 0
                try: rad = int(sh.adjustments[0] * min(w, h))
                except Exception: pass
                if 'OVAL' in str(kind): dr.ellipse([x, y, x + w, y + h], fill=fill, outline=outline, width=3)
                elif rad: dr.rounded_rectangle([x, y, x + w, y + h], radius=rad, fill=fill, outline=outline, width=3)
                else: dr.rectangle([x, y, x + w, y + h], fill=fill, outline=outline)
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            para = tf.paragraphs[0]
            runs = [(r.text, bool(r.font.bold), r.font.name, r.font.size.pt * 2) for r in para.runs]
            ls = para.line_spacing or 1.0
            lines = wrap(runs, w)
            lh = [s * 1.2 * ls for _, s in lines]
            tot = sum(lh)
            col = tuple(para.runs[0].font.color.rgb)
            single = len(lines) == 1
            if tot > h * 1.03 and not (single and abs(tot - h) < 6):
                problems.append(f'Diap {si}: texto desborda su caja ({tot:.0f}px de texto en {h:.0f}px): "{tf.text[:48]}"')
            yy = y
            if tf.vertical_anchor is not None and 'MIDDLE' in str(tf.vertical_anchor): yy = y + (h - tot) / 2
            for (ln, s), hh in zip(lines, lh):
                total_w = sum(ImageDraw.Draw(Image.new('RGB', (1, 1))).textlength(wd + ' ', font=font(n, b, sz)) for wd, b, n, sz in ln)
                xx = x
                if para.alignment is not None and 'CENTER' in str(para.alignment): xx = x + (w - total_w) / 2
                for wd, b, n, sz in ln:
                    f = font(n, b, sz); dr.text((xx, yy), wd + ' ', font=f, fill=col)
                    xx += dr.textlength(wd + ' ', font=f)
                yy += hh
    canvas.resize((sheet_w, sheet_h)).save(os.path.join(os.environ['TEMP'], f'pptx_{si}.png'))
    sheet.paste(canvas.resize((sheet_w, sheet_h)), (((si - 1) % 3) * sheet_w, ((si - 1) // 3) * sheet_h))
    if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
        problems.append(f'Diap {si}: sin notas del orador')
sheet.save(os.path.join(os.environ['TEMP'], 'pptx_sheet.png'))
print(len(prs.slides), 'diapositivas')
print('PROBLEMAS:' if problems else 'Sin problemas detectados')
for p in problems: print(' -', p)
