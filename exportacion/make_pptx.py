from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image
import os, copy

PX = 6350  # 1920 px = 13.333 in
IMG = os.path.join(os.path.dirname(__file__), '..', 'presentacion_img')
NAVY, LIGHT, ORANGE, GREEN, BLUE = '10243A', 'F7F5F0', 'E8791B', '0E9347', '7EC4EE'
HEAD, BODY = 'Arial Narrow', 'Calibri'

prs = Presentation()
prs.slide_width = Emu(1920 * PX)
prs.slide_height = Emu(1080 * PX)
blank = prs.slide_layouts[6]

def E(v): return Emu(int(v * PX))
def rgb(h): return RGBColor.from_string(h)

def set_alpha(fill_elm_parent, alpha):
    clr = fill_elm_parent.find('.//' + qn('a:srgbClr'))
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(alpha * 1000)))

def rect(s, x, y, w, h, fill, alpha=None, radius=None, line=None, shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sp = s.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, radius / min(w, h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None: set_alpha(sp.fill._xPr.find(qn('a:solidFill')), alpha)
    if line:
        sp.line.color.rgb = rgb(line[0]); sp.line.width = Emu(int(line[1] * PX))
    else:
        sp.line.fill.background()
    return sp

def grad(s, x, y, w, h, angle, stops):
    sp = rect(s, x, y, w, h, NAVY)
    spPr = sp._element.spPr
    for tag in ('a:solidFill', 'a:gradFill'):
        for e in spPr.findall(qn(tag)): spPr.remove(e)
    g = etree.Element(qn('a:gradFill')); g.set('rotWithShape', '1')
    gl = etree.SubElement(g, qn('a:gsLst'))
    for pos, col, al in stops:
        gs = etree.SubElement(gl, qn('a:gs')); gs.set('pos', str(int(pos * 1000)))
        c = etree.SubElement(gs, qn('a:srgbClr')); c.set('val', col)
        a = etree.SubElement(c, qn('a:alpha')); a.set('val', str(int(al * 1000)))
    lin = etree.SubElement(g, qn('a:lin')); lin.set('ang', str(int(angle * 60000))); lin.set('scaled', '0')
    spPr.insert(list(spPr).index(spPr.find(qn('a:prstGeom'))) + 1, g)
    return sp

def pic(s, name, x, y, w, h):
    path = os.path.join(IMG, name)
    p = s.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    iw, ih = Image.open(path).size
    ar_box, ar_img = w / h, iw / ih
    if ar_img > ar_box:
        c = (1 - ar_box / ar_img) / 2; p.crop_left = c; p.crop_right = c
    else:
        c = (1 - ar_img / ar_box) / 2; p.crop_top = c; p.crop_bottom = c
    return p

def pic_fit(s, name, x, y, w, h):
    path = os.path.join(IMG, name)
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih); nw, nh = iw * sc, ih * sc
    return s.shapes.add_picture(path, E(x + (w - nw) / 2), E(y + (h - nh) / 2), E(nw), E(nh))

def text(s, x, y, w, h, runs, size, color, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=1.0):
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]; para.alignment = align; para.line_spacing = line
    if isinstance(runs, str): runs = [(runs, bold)]
    for t, b in runs:
        r = para.add_run(); r.text = t
        r.font.size = Pt(size * 0.5); r.font.bold = b; r.font.name = font
        r.font.color.rgb = rgb(color)
        if spacing is not None:
            r._r.get_or_add_rPr().set('spc', str(int(spacing * 50)))
    return tb

def bg(s, color):
    f = s.background.fill; f.solid(); f.fore_color.rgb = rgb(color)

def notes(s, t): s.notes_slide.notes_text_frame.text = t

def new(color, note):
    s = prs.slides.add_slide(blank); bg(s, color); notes(s, note); return s

def logos(s, x, y):
    rect(s, x, y, 584, 166, LIGHT, radius=16)
    pic_fit(s, 'logo_ayto.png', x + 32, y + 20, 200, 126)
    pic_fit(s, 'logo_deporte.png', x + 272, y + 51, 280, 64)

def badge(s, x, y, d, fill, label, color='FFFFFF', size=None):
    c = rect(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    text(s, x, y, d, d, label, size or d * 0.5, color, 'Segoe UI Symbol' if label == '✓' else HEAD, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 1 · Portada ---------------------------------------------------------------
s = new(NAVY, 'Bienvenida. Presentamos el proyecto de mejora de las pistas multideporte del Polideportivo Juan Carlos I: qué problema resuelve, por qué es necesario y cómo será el nuevo módulo.')
pic(s, 'r_ext.jpg', 0, 0, 1920, 1080)
grad(s, 0, 0, 1920, 1080, 0, [(0, NAVY, 94), (42, NAVY, 78), (100, NAVY, 5)])
logos(s, 128, 80)
text(s, 128, 466, 1100, 36, 'POLIDEPORTIVO JUAN CARLOS I · CIUDAD REAL', 28, 'F29A4A', BODY, True, spacing=4)
text(s, 128, 520, 1150, 320, 'MEJORA DE PISTAS MULTIDEPORTE', 136, LIGHT, HEAD, True, line=0.92)
text(s, 128, 868, 960, 110, 'Un espacio cubierto para entrenar, competir y crecer todo el año.', 42, 'E3EAF2')

# 2 · Hoy ------------------------------------------------------------------
s = new(LIGHT, 'Empezamos por la realidad: tenemos un espacio cubierto muy grande, pero que no rinde todo lo que podría. El pavimento está desgastado y la nave está abierta al exterior.')
pic(s, 'antes.jpg', 128, 204, 880, 640)
text(s, 1088, 236, 704, 36, 'PUNTO DE PARTIDA', 28, 'B95A0C', BODY, True, spacing=4)
text(s, 1088, 290, 704, 380, 'UNA GRAN PISTA CUBIERTA QUE PEDÍA RENOVARSE', 84, NAVY, HEAD, True, line=0.95)
text(s, 1088, 690, 704, 260, 'El Polideportivo Juan Carlos I cuenta con un amplio espacio cubierto. Hoy tiene el pavimento desgastado, una nave abierta a la lluvia, al viento y a las aves, y una sola pista para todos los deportes.', 36, '3E4C5E', line=1.05)
text(s, 128, 980, 1200, 30, 'Estado actual de la pista · Polideportivo Juan Carlos I', 24, '6B7888')

# 3 · Necesidad ------------------------------------------------------------
s = new(NAVY, 'Tres motivos claros. Uno, el pavimento sufre. Dos, la nave no protege ni del tiempo ni de las aves. Tres, no hay un espacio cubierto específico para el atletismo ni una forma ordenada de repartir los usos.')
text(s, 128, 128, 1000, 36, 'POR QUÉ ES NECESARIO', 28, 'F29A4A', BODY, True, spacing=4)
text(s, 128, 176, 1664, 120, 'TRES RAZONES PARA ACTUAR', 96, LIGHT, HEAD, True)
cards = [('1', 'PAVIMENTO DETERIORADO', 'El uso continuado y las aves que anidan dentro dañan la superficie de juego.'),
         ('2', 'UNA NAVE A LA INTEMPERIE', 'Lluvia, frío, viento y suciedad entran con facilidad y limitan el uso del espacio.'),
         ('3', 'UN ESPACIO PARA TODO', 'No hay un lugar cubierto para el atletismo, ni se pueden separar los usos de forma ordenada.')]
cw = (1664 - 80) / 3
for i, (n, t, d) in enumerate(cards):
    x = 128 + i * (cw + 40)
    rect(s, x, 410, cw, 500, '1B3652', radius=20)
    rect(s, x + 24, 410, cw - 48, 8, ORANGE)
    badge(s, x + 48, 470, 84, ORANGE, n, NAVY, 52)
    text(s, x + 48, 580, cw - 96, 130, t, 52, LIGHT, HEAD, True, line=0.95)
    text(s, x + 48, 730, cw - 96, 160, d, 32, 'D5E0EC', line=1.05)

# 4 · Antes / después -----------------------------------------------------
s = new(LIGHT, 'Aquí se ve el salto: a la izquierda, cómo está hoy; a la derecha, una recreación en 3D del módulo ya renovado, con pistas multideporte y atletismo bajo la misma cubierta.')
text(s, 128, 96, 1664, 190, 'DE UN ESPACIO ABIERTO A UN MÓDULO DEPORTIVO COMPLETO', 80, NAVY, HEAD, True, line=0.95)
pic(s, 'antes.jpg', 128, 330, 812, 520)
pic(s, 'r_int.jpg', 980, 330, 812, 520)
text(s, 128, 876, 812, 40, 'HOY', 30, '3E4C5E', BODY, True, spacing=3)
text(s, 980, 876, 812, 40, 'EL MÓDULO', 30, 'B95A0C', BODY, True, spacing=3)

# 5 · Respuesta -----------------------------------------------------------
s = new(NAVY, 'La respuesta es sencilla de explicar: aprovechar el mismo espacio para tres cosas. Dos pistas multideporte, un módulo de atletismo cubierto y una nave renovada que lo protege todo.')
text(s, 128, 118, 800, 36, 'LA RESPUESTA', 28, 'F29A4A', BODY, True, spacing=4)
text(s, 128, 170, 800, 520, 'DOS PISTAS Y UN MÓDULO DE ATLETISMO, BAJO UNA MISMA CUBIERTA', 88, LIGHT, HEAD, True, line=0.95)
for i, (t, w, f, c) in enumerate([('2 pistas multideporte', 420, ORANGE, NAVY), ('1 módulo de atletismo cubierto', 560, GREEN, LIGHT), ('1 nave renovada y protegida', 520, BLUE, NAVY)]):
    y = 720 + i * 82
    rect(s, 128, y, w, 66, f, radius=12)
    text(s, 128, y, w, 66, t, 34, c, BODY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
pic(s, 'r_atl2.jpg', 1000, 190, 820, 700)

# 5b · Planta (vista desde arriba) ----------------------------------------
s = new(LIGHT, 'Así se organiza el módulo visto desde arriba: arriba, las ocho calles de atletismo; en el centro, las zonas de salto; y abajo, las dos pistas multideporte, separadas entre sí por cortinas.')
pic(s, 'r_planta.jpg', 240, 270, 1440, 720)
text(s, 128, 80, 1664, 100, 'TODO EN UN VISTAZO', 80, NAVY, HEAD, True)
x = 128
for n, lab, w in [('1', 'Atletismo · 8 calles', 380), ('2', 'Zonas de salto', 300), ('3', 'Pista multideporte 1', 380), ('4', 'Pista multideporte 2', 380)]:
    badge(s, x, 196, 44, ORANGE, n, NAVY, 28)
    text(s, x + 60, 196, w, 44, lab, 28, NAVY, BODY, True, anchor=MSO_ANCHOR.MIDDLE)
    x += 60 + w + 32
for n, cx, cy in [('1', 1126, 464), ('2', 1032, 572), ('3', 686, 738), ('4', 1234, 738)]:
    badge(s, cx - 32, cy - 32, 64, ORANGE, n, NAVY, 36)

# 6 · Atletismo -----------------------------------------------------------
s = new(NAVY, 'Este es el corazón del proyecto: un módulo de atletismo cubierto, con ocho calles de noventa metros, zona de frenada incluida, y zonas de salto. Se entrega homologado por la Real Federación Española de Atletismo. Permite entrenar y tecnificar todo el año, llueva o haga frío.')
pic(s, 'r_atl.jpg', 0, 0, 1920, 1080)
grad(s, 0, 0, 1920, 1080, 90, [(0, NAVY, 90), (38, NAVY, 25), (60, NAVY, 35), (100, NAVY, 94)])
text(s, 128, 96, 1000, 36, 'EL MÓDULO DE ATLETISMO', 28, 'F29A4A', BODY, True, spacing=4)
text(s, 128, 144, 1400, 250, 'ATLETISMO, BAJO CUBIERTA', 104, LIGHT, HEAD, True, line=0.95)
stats = [('8', 'calles de velocidad', ORANGE), ('90 m', 'de pista de velocidad, con zona de frenada', ORANGE),
         ('4', 'zonas de salto: altura, longitud, triple y pértiga', ORANGE), ('RFEA', 'módulo homologado por la federación', GREEN)]
for i, (n, l, c) in enumerate(stats):
    x = 128 + i * (380 + 48)
    rect(s, x, 640, 380, 6, c)
    text(s, x, 664, 380, 140, n, 128, LIGHT, HEAD, True, line=0.9)
    text(s, x, 810, 380, 120, l, 32, 'E3EAF2', line=1.05)

# 7 · Pistas ---------------------------------------------------------------
s = new(LIGHT, 'Las dos pistas multideporte admiten fútbol sala, baloncesto, voleibol y tenis. Y gracias a las cortinas motorizadas, se pueden usar a la vez que el atletismo sin interferirse.')
text(s, 128, 214, 720, 36, 'DOS PISTAS MULTIDEPORTE', 28, 'B95A0C', BODY, True, spacing=4)
text(s, 128, 264, 720, 320, 'UNA PISTA, MUCHOS DEPORTES', 96, NAVY, HEAD, True, line=0.95)
for t, x, y, w in [('Fútbol sala', 128, 610, 220), ('Baloncesto', 364, 610, 220), ('Voleibol', 600, 610, 190), ('Tenis', 128, 684, 140)]:
    rect(s, x, y, w, 58, None, radius=29, line=(ORANGE, 3))
    text(s, x, y, w, 58, t, 30, NAVY, BODY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 128, 780, 720, 200, 'Cortinas motorizadas separan las pistas entre sí y del atletismo: varios grupos entrenan a la vez, sin molestarse.', 34, '3E4C5E', line=1.05)
pic(s, 'r_cancha.jpg', 920, 190, 900, 700)

# 7b · Superficies y fases ------------------------------------------------
s = new(LIGHT, 'Cada zona lleva su propia superficie. En atletismo, un tartán sintético poroso, aplicado in situ y homologado por la RFEA. En las pistas multideporte, un pavimento de PVC o resinas, como el de los pabellones cubiertos. Como las superficies son independientes, la obra puede hacerse por fases y adaptar el orden. El proyecto estima unos seis meses en total.')
text(s, 128, 72, 1000, 36, 'LAS SUPERFICIES', 28, 'B95A0C', BODY, True, spacing=4)
text(s, 128, 112, 1664, 100, 'DOS SUPERFICIES, CADA UNA PARA LO SUYO', 80, NAVY, HEAD, True)
cards = [(128, 'sup_tartan.jpg', 'ATLETISMO', GREEN, LIGHT, 200, 'TARTÁN SINTÉTICO',
          ['Poroso y aplicado in situ', 'Acabado en azul con resina roja', 'Homologado por la RFEA']),
         (976, 'sup_pvc.jpg', 'PISTAS MULTIDEPORTE', ORANGE, NAVY, 420, 'PAVIMENTO DE PVC O RESINAS',
          ['Como en los pabellones cubiertos', 'Para fútbol sala, baloncesto, voleibol y tenis', 'Líneas de cada deporte en la misma pista'])]
for x, img, tag, fc, tc, tw, title, buls in cards:
    rect(s, x, 230, 816, 575, 'FFFFFF', radius=20)
    pic(s, img, x, 230, 816, 250)
    rect(s, x + 24, 795, 768, 10, fc)
    rect(s, x + 36, 510, tw, 44, fc, radius=22)
    text(s, x + 36, 510, tw, 44, tag, 26, tc, BODY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, spacing=3)
    text(s, x + 36, 568, 744, 62, title, 50, NAVY, HEAD, True)
    for i, b in enumerate(buls):
        text(s, x + 36, 642 + i * 50, 744, 44, '•  ' + b, 30, '3E4C5E')
text(s, 128, 833, 1664, 50, 'SE PUEDE EJECUTAR POR FASES', 40, NAVY, HEAD, True, spacing=2)
for i, lab in enumerate(['1 · Módulo de atletismo', '2 · Pistas multideporte', '3 · Cubierta y cerramiento']):
    cx = 128 + i * (517 + 56)
    rect(s, cx, 896, 517, 66, NAVY, radius=14)
    text(s, cx, 896, 517, 66, lab, 28, LIGHT, BODY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 2:
        text(s, cx + 517, 896, 56, 66, '→', 40, ORANGE, BODY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 128, 980, 1664, 40, 'Cada superficie es independiente, así que el orden se puede adaptar. Estimación del proyecto: unos 6 meses en total.', 26, '3E4C5E')

# 8 · Nave -----------------------------------------------------------------
s = new(NAVY, 'La nave también se renueva: nueva cubierta aislante, un cerramiento que deja pasar la luz, y un entorno exterior más cómodo y accesible. El resultado es un espacio protegido de la lluvia, el frío, el viento y las aves.')
pic(s, 'r_ext.jpg', 128, 180, 860, 720)
text(s, 1060, 150, 732, 290, 'UNA NAVE PROTEGIDA Y CONFORTABLE', 84, LIGHT, HEAD, True, line=0.95)
rows = [('Cubierta renovada.', ' Más aislamiento y menos frío en invierno.'),
        ('Luz natural.', ' Cerramiento translúcido y control del sol en la fachada oeste.'),
        ('Espacio protegido.', ' Sin aves, sin lluvia y sin suciedad sobre la pista.'),
        ('Entorno accesible.', ' Mejor recogida de agua y acceso peatonal desde el exterior.')]
for i, (b, t) in enumerate(rows):
    y = 470 + i * 106
    badge(s, 1060, y + 4, 56, ORANGE, '✓', NAVY, 30)
    text(s, 1140, y, 652, 100, [(b, True), (t, False)], 32, LIGHT, line=1.05)

# 9 · Beneficios -----------------------------------------------------------
s = new(LIGHT, 'Los beneficiarios son muy distintos: clubes y escuelas, atletas que quieren tecnificarse, organizadores de competiciones y cualquier vecino que quiera hacer deporte.')
text(s, 128, 128, 1000, 36, 'QUÉ GANA CIUDAD REAL', 28, 'B95A0C', BODY, True, spacing=4)
text(s, 128, 176, 1664, 120, 'UN ESPACIO PARA TODOS', 96, NAVY, HEAD, True)
tiles = [('CLUBES Y ESCUELAS', 'Un espacio cubierto y versátil para entrenar con cualquier clima.', ORANGE, '1'),
         ('ATLETAS EN TECNIFICACIÓN', 'Calles y zonas de salto para trabajar la técnica durante todo el año.', GREEN, '2'),
         ('COMPETICIONES Y EVENTOS', 'Módulo de atletismo homologado por la RFEA.', GREEN, '3'),
         ('DEPORTE PARA TODOS', 'Dos pistas multideporte abiertas a deportistas de todas las edades.', ORANGE, '4')]
for i, (t, d, c, n) in enumerate(tiles):
    x = 128 + (i % 2) * (816 + 32); y = 410 + (i // 2) * (260 + 32)
    rect(s, x, y, 816, 260, 'FFFFFF', radius=20)
    rect(s, x, y + 20, 10, 220, c)
    badge(s, x + 44, y + 44, 72, c, n, 'FFFFFF', 40)
    text(s, x + 144, y + 40, 640, 60, t, 46, NAVY, HEAD, True)
    text(s, x + 144, y + 108, 640, 130, d, 30, '3E4C5E', line=1.05)

# 10 · Cierre --------------------------------------------------------------
s = new(NAVY, 'Cerramos con la idea central: un espacio cubierto que permite practicar atletismo y deporte todo el año. Gracias.')
pic(s, 'r_int.jpg', 0, 0, 1920, 1080)
grad(s, 0, 0, 1920, 1080, 0, [(0, NAVY, 95), (50, NAVY, 80), (100, NAVY, 30)])
text(s, 128, 310, 1000, 36, 'UN MÓDULO PARA CIUDAD REAL', 28, 'F29A4A', BODY, True, spacing=4)
text(s, 128, 360, 1400, 240, 'ATLETISMO Y DEPORTE BAJO CUBIERTA, TODO EL AÑO', 104, LIGHT, HEAD, True, line=0.95)
text(s, 128, 620, 1000, 60, 'Un espacio para entrenar, tecnificar y competir.', 40, 'E3EAF2')
logos(s, 128, 834)

out = os.path.join(os.path.dirname(__file__), 'Presentacion_Mejora_Pistas_Multideporte.pptx')
prs.save(out); print('guardado', out)
